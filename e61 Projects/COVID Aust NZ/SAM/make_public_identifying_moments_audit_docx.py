"""Build the Word version of the public identifying-moments audit.

The substantive source is Public_identifying_moments_audit.md.  Keeping the
source and renderer beside the model makes the report easy to update when
secure-environment moments or public-data vintages change.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Public_identifying_moments_audit.md"
OUTPUT = HERE / "Public_Identifying_Moments_Audit_and_Model16_Next_Steps.docx"

TEAL = "007A78"
DARK_TEAL = "005B5A"
LIGHT_TEAL = "E7F3F2"
MID_GREY = "D9E1E3"
LIGHT_GREY = "F4F6F7"
DARK_GREY = "3F4A4C"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start),
                          ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    relation_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    new_run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    run_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(underline)
    new_run.append(run_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


INLINE_PATTERN = re.compile(
    r"(\[([^\]]+)\]\((https?://[^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`|\*([^*]+)\*)"
)


def add_inline(paragraph, text: str) -> None:
    """Render the small Markdown subset used in the audit."""
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()])
        token = match.group(0)
        if token.startswith("["):
            add_hyperlink(paragraph, match.group(2), match.group(3))
        elif token.startswith("**"):
            run = paragraph.add_run(match.group(4))
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(match.group(5))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(DARK_TEAL)
        elif token.startswith("*"):
            run = paragraph.add_run(match.group(6))
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Open in Word and update this field to populate the table of contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor.from_string(DARK_GREY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK_TEAL)
    title.paragraph_format.space_after = Pt(16)

    for name, size, before, after, color in (
        ("Heading 1", 16, 14, 7, DARK_TEAL),
        ("Heading 2", 13, 11, 5, TEAL),
        ("Heading 3", 11, 8, 3, DARK_GREY),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style.font.size = Pt(10.3)
        style.paragraph_format.space_after = Pt(3)

    if "Audit Quote" not in styles:
        quote = styles.add_style("Audit Quote", WD_STYLE_TYPE.PARAGRAPH)
    else:
        quote = styles["Audit Quote"]
    quote.base_style = styles["Normal"]
    quote.font.name = "Aptos"
    quote.font.size = Pt(11)
    quote.font.bold = True
    quote.font.color.rgb = RGBColor.from_string(DARK_TEAL)
    quote.paragraph_format.left_indent = Cm(0.7)
    quote.paragraph_format.right_indent = Cm(0.4)
    quote.paragraph_format.space_before = Pt(7)
    quote.paragraph_format.space_after = Pt(9)

    doc.core_properties.title = (
        "Public identifying moments audit and recommended next steps for model 16"
    )
    doc.core_properties.subject = (
        "External data audit and structural identification plan for the COVID "
        "Australia-New Zealand labour-supply paper"
    )
    doc.core_properties.author = "Prepared for the research team"
    doc.core_properties.keywords = (
        "JobSeeker, Coronavirus Supplement, JobKeeper, vacancies, search and matching, identification"
    )

    header = section.header.paragraphs[0]
    header.text = "PUBLIC IDENTIFYING MOMENTS AUDIT  |  MODEL 16"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(TEAL)

    footer = section.footer.paragraphs[0]
    footer.add_run("COVID labour-supply structural model  |  ")
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(DARK_GREY)


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator_row(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    n_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        for c_idx in range(n_cols):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            value = values[c_idx] if c_idx < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Aptos"
                run.font.size = Pt(8.2)
                if r_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            if r_idx == 0:
                set_cell_shading(cell, TEAL)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_GREY)
    set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_paragraph(doc: Document, text: str, style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    add_inline(paragraph, text)
    return paragraph


def render_markdown(doc: Document, lines: list[str]) -> None:
    index = 0
    title_seen = False
    toc_inserted = False
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_paragraph(doc, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            index += 1
            continue

        if stripped.startswith("#"):
            flush_paragraph()
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and not title_seen:
                paragraph = doc.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline(paragraph, text)
                title_seen = True
            else:
                word_level = min(max(level - 1, 1), 3)
                heading = doc.add_heading(text, level=word_level)
                if word_level == 1 and (
                    re.match(r"^\d+\.", text) or text.startswith("Appendix")
                ):
                    heading.paragraph_format.page_break_before = True
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and is_separator_row(lines[index + 1]):
            flush_paragraph()
            table_rows = [split_table_row(stripped)]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_rows.append(split_table_row(lines[index]))
                index += 1
            add_table(doc, table_rows)
            continue

        if stripped.startswith(">"):
            flush_paragraph()
            quote_parts = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_parts.append(lines[index].strip()[1:].strip())
                index += 1
            paragraph = add_paragraph(doc, " ".join(quote_parts), "Audit Quote")
            p_pr = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), TEAL)
            borders.append(left)
            p_pr.append(borders)
            continue

        bullet_match = re.match(r"^-\s+(.*)$", stripped)
        number_match = re.match(r"^\d+\.\s+(.*)$", stripped)
        if bullet_match:
            flush_paragraph()
            add_paragraph(doc, bullet_match.group(1), "List Bullet")
            index += 1
            continue
        if number_match:
            flush_paragraph()
            add_paragraph(doc, number_match.group(1), "List Number")
            index += 1
            continue

        # The first metadata block doubles as the title-page subtitle. Insert a
        # field-based table of contents immediately before the executive summary.
        if stripped.startswith("**") and not toc_inserted and title_seen:
            flush_paragraph()
            paragraph = add_paragraph(doc, stripped)
            paragraph.paragraph_format.space_after = Pt(2)
            index += 1
            if index < len(lines) and lines[index].strip() == "":
                # The block is handled line-by-line, so wait until Purpose below.
                pass
            if stripped.startswith("**Purpose:"):
                doc.add_paragraph()
                toc_heading = doc.add_heading("Contents", level=1)
                toc_heading.paragraph_format.page_break_before = True
                toc = doc.add_paragraph()
                add_toc(toc)
                doc.add_page_break()
                toc_inserted = True
            continue

        paragraph_buffer.append(stripped)
        index += 1

    flush_paragraph()


def improve_pagination(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name in {
            "Heading 1", "Heading 2", "Heading 3"
        }:
            paragraph.paragraph_format.keep_with_next = True
        if paragraph.style and paragraph.style.name in {"List Bullet", "List Number"}:
            paragraph.paragraph_format.widow_control = True


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document()
    configure_document(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    render_markdown(doc, lines)
    improve_pagination(doc)
    doc.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    main()
